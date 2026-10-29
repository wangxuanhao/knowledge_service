"""Native Docling media survives normalization without embedded payloads."""
import io
from pathlib import Path
import pytest
from PIL import Image
from knowledge_service.integrations.document_parser import DocumentParser, DocumentParseError


def test_native_document_keeps_picture_positions_bytes_and_provenance(tmp_path):
    from docling_core.types.doc import DoclingDocument, DocItemLabel, ImageRef
    document=DoclingDocument(name='article')
    document.add_text(label=DocItemLabel.PARAGRAPH,text='Before picture')
    document.add_picture(image=ImageRef.from_pil(Image.new('RGB',(12,8),'red'),dpi=72))
    document.add_text(label=DocItemLabel.PARAGRAPH,text='Between pictures')
    document.add_picture(image=ImageRef.from_pil(Image.new('RGB',(12,8),'blue'),dpi=72))
    document.add_text(label=DocItemLabel.PARAGRAPH,text='After picture')
    class Parser:
        def __init__(self,**kwargs): pass
        def parse(self,path): return {'full_text':'WRAPPER DROPPED MEDIA','document':document}
    path=tmp_path/'a.docx';path.write_bytes(b'fixture')
    parsed=DocumentParser(parser_factory=Parser).parse(path,'a.docx')
    assert len(getattr(parsed,'assets',[]))==2
    assert parsed.text.index('Before')<parsed.text.index('assets/picture-1.png')<parsed.text.index('Between')<parsed.text.index('assets/picture-2.png')<parsed.text.index('After')
    assert 'base64' not in parsed.text
    assert parsed.assets[0]['data'].startswith(b'\x89PNG')
    assert len(parsed.provenance['items'])==5


def test_real_docx_keeps_images_in_body_order_without_pdf_models(tmp_path,monkeypatch):
    from docx import Document
    image=io.BytesIO();Image.new('RGB',(12,8),'red').save(image,format='PNG');image.seek(0)
    document=Document();document.add_paragraph('Before picture');document.add_picture(image);document.add_paragraph('After picture')
    path=tmp_path/'article.docx';document.save(path)
    monkeypatch.setenv('KG_DOCUMENT_ARTIFACTS_PATH',str(tmp_path/'missing'))
    monkeypatch.setattr(DocumentParser,'_factory',lambda self:pytest.fail('Native Docling must not initialize the lossy Semantica wrapper'))
    parsed=DocumentParser(ocr_mode='disabled').parse(path,path.name)
    assert len(getattr(parsed,'assets',[]))==1
    assert parsed.text.index('Before')<parsed.text.index('assets/picture-1.png')<parsed.text.index('After')


def test_missing_pdf_models_fail_before_conversion(tmp_path,monkeypatch):
    monkeypatch.setenv('KG_DOCUMENT_ARTIFACTS_PATH',str(tmp_path/'missing'))
    path=tmp_path/'a.pdf';path.write_bytes(b'%PDF')
    with pytest.raises(DocumentParseError) as error: DocumentParser(ocr_mode='disabled').parse(path,path.name)
    assert error.value.code=='model_unavailable'


def test_pdf_options_use_explicit_local_models_and_really_disable_ocr(tmp_path,monkeypatch):
    from docling.datamodel.base_models import InputFormat
    # This test constructs options, without running model inference. Supply the
    # preflight files explicitly so a clean checkout needs no downloaded weights.
    for name in ['docling-project--docling-layout-heron/config.json',
                 'docling-project--docling-layout-heron/preprocessor_config.json',
                 'docling-project--docling-layout-heron/model.safetensors',
                 'docling-project--docling-models/model_artifacts/tableformer/accurate/tm_config.json',
                 'docling-project--docling-models/model_artifacts/tableformer/accurate/tableformer_accurate.safetensors']:
        artifact=tmp_path/name;artifact.parent.mkdir(parents=True,exist_ok=True);artifact.write_bytes(b'fixture')
    monkeypatch.setenv('KG_DOCUMENT_ARTIFACTS_PATH',str(tmp_path))
    parser=DocumentParser(ocr_mode='disabled')._new_docling_parser('.pdf')
    options=parser._converter.format_to_options[InputFormat.PDF].pipeline_options
    assert not options.do_ocr and options.generate_picture_images
    assert options.artifacts_path.is_absolute()
    assert options.artifacts_path==tmp_path


def test_scanned_page_fallback_preserves_page_image_without_adding_digital_page_screenshots():
    from types import SimpleNamespace
    from docling_core.types.doc import DoclingDocument,DocItemLabel,ImageRef,Size,ProvenanceItem,BoundingBox
    from knowledge_service.integrations.document_media import normalize_docling_document
    document=DoclingDocument(name='mixed')
    for number,text in [(1,'Digital article'),(2,'OCR reading of scanned page')]:
        document.add_page(page_no=number,size=Size(width=100,height=80),image=ImageRef.from_pil(Image.new('RGB',(100,80),'white'),dpi=72))
        document.add_text(label=DocItemLabel.PARAGRAPH,text=text,prov=ProvenanceItem(page_no=number,bbox=BoundingBox(l=0,t=0,r=90,b=60),charspan=(0,len(text))))
    text,assets,provenance,warnings=normalize_docling_document(document,scanned_pages=[2])
    assert len(assets)==1 and assets[0]['provenance']['page_no']==2
    assert 'assets/page-2.png' in text
    assert text.index('Digital article')<text.index('assets/page-2.png')
    assert 'OCR reading' not in text


def test_empty_scanned_page_is_readable_with_ocr_disabled():
    from docling_core.types.doc import DoclingDocument,ImageRef,Size
    from knowledge_service.integrations.document_media import normalize_docling_document
    document=DoclingDocument(name='scan')
    document.add_page(page_no=1,size=Size(width=100,height=80),image=ImageRef.from_pil(Image.new('RGB',(100,80),'white'),dpi=72))
    text,assets,_,_=normalize_docling_document(document,scanned_pages=[1])
    assert text.strip() and len(assets)==1


def test_native_ocr_text_is_separated_from_digital_text_on_mixed_pages():
    from docling_core.types.doc import DoclingDocument,DocItemLabel,Size,ProvenanceItem,BoundingBox,CoordOrigin
    from knowledge_service.integrations.document_media import normalize_docling_document
    document=DoclingDocument(name='mixed')
    document.add_page(page_no=1,size=Size(width=100,height=100))
    def box(top,bottom):return BoundingBox(l=0,t=top,r=100,b=bottom,coord_origin=CoordOrigin.TOPLEFT)
    document.add_text(label=DocItemLabel.PARAGRAPH,text='APPROVED 42 digital source',prov=ProvenanceItem(page_no=1,bbox=box(0,20),charspan=(0,26)))
    document.add_text(label=DocItemLabel.PARAGRAPH,text='APPROVED 42',prov=ProvenanceItem(page_no=1,bbox=box(50,70),charspan=(0,11)))
    text,_,_,_=normalize_docling_document(document,ocr_regions=[dict(page_no=1,text='APPROVED 42',bbox=box(50,70).model_dump(mode='json'))])
    assert text.split('![')[0].strip()=='APPROVED 42 digital source'


def test_native_table_value_is_not_erased_by_ocr_text_in_another_cell():
    from docling_core.types.doc import DoclingDocument,TableData,TableCell,Size,ProvenanceItem,BoundingBox,CoordOrigin
    from knowledge_service.integrations.document_media import normalize_docling_document
    document=DoclingDocument(name='table')
    document.add_page(page_no=1,size=Size(width=100,height=100))
    def box(left,right):return BoundingBox(l=left,t=10,r=right,b=40,coord_origin=CoordOrigin.TOPLEFT)
    document.add_table(data=TableData(num_rows=1,num_cols=2,table_cells=[
        TableCell(text='42',start_row_offset_idx=0,end_row_offset_idx=1,start_col_offset_idx=0,end_col_offset_idx=1,bbox=box(0,45)),
        TableCell(text='42',start_row_offset_idx=0,end_row_offset_idx=1,start_col_offset_idx=1,end_col_offset_idx=2,bbox=box(55,100))]),
        prov=ProvenanceItem(page_no=1,bbox=box(0,100),charspan=(0,4)))
    text,_,_,_=normalize_docling_document(document,ocr_regions=[dict(page_no=1,text='42',bbox=box(55,100).model_dump(mode='json'))])
    assert text.count('42')==1


def test_native_value_in_shared_text_item_survives_ocr_removal_in_other_charspan():
    from docling_core.types.doc import DoclingDocument,DocItemLabel,Size,ProvenanceItem,BoundingBox,CoordOrigin
    from knowledge_service.integrations.document_media import normalize_docling_document
    document=DoclingDocument(name='shared text')
    document.add_page(page_no=1,size=Size(width=100,height=100))
    def box(left,right):return BoundingBox(l=left,t=0,r=right,b=20,coord_origin=CoordOrigin.TOPLEFT)
    item=document.add_text(label=DocItemLabel.PARAGRAPH,text='42 42',prov=ProvenanceItem(page_no=1,bbox=box(0,40),charspan=(0,2)))
    item.prov.append(ProvenanceItem(page_no=1,bbox=box(60,100),charspan=(3,5)))
    text,_,_,_=normalize_docling_document(document,ocr_regions=[dict(page_no=1,text='42',bbox=box(60,100).model_dump(mode='json'))])
    assert text.split('![')[0].strip()=='42'


def test_normalized_native_ocr_text_cannot_bypass_common_quality_gate():
    from docling_core.types.doc import DoclingDocument,DocItemLabel,Size,ProvenanceItem,BoundingBox,CoordOrigin
    from knowledge_service.integrations.document_media import normalize_docling_document
    document=DoclingDocument(name='normalized OCR')
    document.add_page(page_no=1,size=Size(width=100,height=100))
    box=BoundingBox(l=0,t=50,r=100,b=80,coord_origin=CoordOrigin.TOPLEFT)
    document.add_text(label=DocItemLabel.PARAGRAPH,text="We're approved",prov=ProvenanceItem(page_no=1,bbox=box,charspan=(0,14)))
    text,_,_,_=normalize_docling_document(document,ocr_regions=[dict(page_no=1,text='We’re approved',bbox=box.model_dump(mode='json'))])
    assert "We're approved" not in text


def test_mixed_digital_header_and_raster_table_retains_table_pixels_in_order():
    from docling_core.types.doc import DoclingDocument,DocItemLabel,TableData,TableCell,Size,ProvenanceItem,BoundingBox,CoordOrigin,ImageRef
    from knowledge_service.integrations.document_media import normalize_docling_document
    document=DoclingDocument(name='raster table')
    document.add_page(page_no=1,size=Size(width=100,height=100))
    header=BoundingBox(l=0,t=0,r=100,b=20,coord_origin=CoordOrigin.TOPLEFT)
    box=BoundingBox(l=0,t=50,r=100,b=80,coord_origin=CoordOrigin.TOPLEFT)
    document.add_text(label=DocItemLabel.PARAGRAPH,text='Native digital header',prov=ProvenanceItem(page_no=1,bbox=header,charspan=(0,21)))
    table=document.add_table(data=TableData(num_rows=1,num_cols=1,table_cells=[TableCell(text='42',start_row_offset_idx=0,end_row_offset_idx=1,start_col_offset_idx=0,end_col_offset_idx=1,bbox=box)]),prov=ProvenanceItem(page_no=1,bbox=box,charspan=(0,2)))
    table.image=ImageRef.from_pil(Image.new('RGB',(100,30),'red'),dpi=72)
    text,assets,_,_=normalize_docling_document(document,ocr_regions=[dict(page_no=1,text='42',bbox=box.model_dump(mode='json'))])
    assert len(assets)==1 and assets[0]['data'].startswith(b'\x89PNG')
    assert text.index('Native digital header')<text.index('assets/table-1.png')
    assert '42' not in text


def test_ambiguous_normalized_span_preserves_only_attributed_digital_cells():
    from docling_core.types.doc import DoclingDocument,DocItemLabel,Size,ProvenanceItem,BoundingBox,CoordOrigin
    from knowledge_service.integrations.document_media import normalize_docling_document
    document=DoclingDocument(name='mixed normalized span')
    document.add_page(page_no=1,size=Size(width=100,height=100))
    def box(left,right):return BoundingBox(l=left,t=0,r=right,b=20,coord_origin=CoordOrigin.TOPLEFT)
    document.add_text(label=DocItemLabel.PARAGRAPH,text="We're approved 42 42",prov=ProvenanceItem(page_no=1,bbox=box(0,100),charspan=(0,21)))
    text,_,_,_=normalize_docling_document(document,
        ocr_regions=[dict(page_no=1,text='We’re approved 42',bbox=box(0,70).model_dump(mode='json'))],
        digital_regions=[dict(page_no=1,text='42',bbox=box(80,100).model_dump(mode='json'))])
    assert text.split('![')[0].strip()=='42'


@pytest.mark.parametrize('spans',[[(0,5),(6,12)],[(0,5),(4,10)],[(0,4),(6,10)]])
def test_inconsistent_dehyphenated_spans_rebuild_whole_item_from_native_cells(spans):
    from docling_core.types.doc import DoclingDocument,DocItemLabel,Size,ProvenanceItem,BoundingBox,CoordOrigin
    from knowledge_service.integrations.document_media import normalize_docling_document
    document=DoclingDocument(name='dehyphenated')
    document.add_page(page_no=1,size=Size(width=100,height=100))
    def box(left,right):return BoundingBox(l=left,t=0,r=right,b=20,coord_origin=CoordOrigin.TOPLEFT)
    item=document.add_text(label=DocItemLabel.PARAGRAPH,text='algorithms',prov=ProvenanceItem(page_no=1,bbox=box(0,40),charspan=spans[0]))
    item.prov.append(ProvenanceItem(page_no=1,bbox=box(60,100),charspan=spans[1]))
    text,_,_,_=normalize_docling_document(document,
        ocr_regions=[dict(page_no=1,text='rithms',bbox=box(60,100).model_dump(mode='json'))],
        digital_regions=[dict(page_no=1,text='algo-',bbox=box(0,40).model_dump(mode='json'))])
    assert text.split('![')[0].strip()=='algo-'


def test_mixed_native_header_and_raster_paragraph_keep_paragraph_pixels_in_order():
    from PIL import ImageDraw
    from docling_core.types.doc import DoclingDocument,DocItemLabel,Size,ProvenanceItem,BoundingBox,CoordOrigin,ImageRef
    from knowledge_service.integrations.document_media import normalize_docling_document
    document=DoclingDocument(name='raster paragraph')
    page=Image.new('RGB',(160,100),'white');ImageDraw.Draw(page).text((0,40),'APPROVED 42',fill='black')
    document.add_page(page_no=1,size=Size(width=160,height=100),image=ImageRef.from_pil(page,dpi=72))
    header=BoundingBox(l=0,t=0,r=160,b=20,coord_origin=CoordOrigin.TOPLEFT)
    raster=BoundingBox(l=0,t=35,r=160,b=70,coord_origin=CoordOrigin.TOPLEFT)
    document.add_text(label=DocItemLabel.PARAGRAPH,text='Native header 42',prov=ProvenanceItem(page_no=1,bbox=header,charspan=(0,16)))
    document.add_text(label=DocItemLabel.PARAGRAPH,text='APPROVED 42',prov=ProvenanceItem(page_no=1,bbox=raster,charspan=(0,11)))
    text,assets,_,_=normalize_docling_document(document,
        ocr_regions=[dict(page_no=1,text='APPROVED 42',bbox=raster.model_dump(mode='json'))],
        digital_regions=[dict(page_no=1,text='Native header 42',bbox=header.model_dump(mode='json'))])
    assert len(assets)==1 and assets[0]['data'].startswith(b'\x89PNG')
    assert text.index('Native header 42')<text.index(assets[0]['original_path'])
    assert 'APPROVED' not in text
    assert assets[0]['width']==160 and assets[0]['height']==35


def test_existing_picture_coverage_avoids_duplicate_ocr_text_crop():
    from docling_core.types.doc import DoclingDocument,DocItemLabel,Size,ProvenanceItem,BoundingBox,CoordOrigin,ImageRef
    from knowledge_service.integrations.document_media import normalize_docling_document
    document=DoclingDocument(name='covered paragraph')
    document.add_page(page_no=1,size=Size(width=100,height=100),image=ImageRef.from_pil(Image.new('RGB',(100,100),'white'),dpi=72))
    box=BoundingBox(l=0,t=30,r=100,b=70,coord_origin=CoordOrigin.TOPLEFT)
    document.add_picture(image=ImageRef.from_pil(Image.new('RGB',(100,40),'red'),dpi=72),prov=ProvenanceItem(page_no=1,bbox=box,charspan=(0,0)))
    document.add_text(label=DocItemLabel.PARAGRAPH,text='APPROVED 42',prov=ProvenanceItem(page_no=1,bbox=box,charspan=(0,11)))
    text,assets,_,_=normalize_docling_document(document,ocr_regions=[dict(page_no=1,text='APPROVED 42',bbox=box.model_dump(mode='json'))])
    assert len(assets)==1 and assets[0]['original_path']=='assets/picture-1.png'
    assert 'APPROVED' not in text


def test_uncaptioned_native_image_does_not_invent_searchable_image_word():
    from docling_core.types.doc import DoclingDocument,ImageRef,DocItemLabel
    from knowledge_service.integrations.document_media import normalize_docling_document
    from knowledge_service.integrations.markdown_parser import parse_markdown
    document=DoclingDocument(name='image only')
    document.add_picture(image=ImageRef.from_pil(Image.new('RGB',(20,20),'red'),dpi=72))
    text,assets,_,_=normalize_docling_document(document)
    assert assets and text.strip()
    assert not ''.join(block['normalized_text'] for block in parse_markdown(text)['blocks'])
    captioned=DoclingDocument(name='captioned image')
    caption=captioned.add_text(label=DocItemLabel.CAPTION,text='Actual caption')
    captioned.add_picture(image=ImageRef.from_pil(Image.new('RGB',(20,20),'red'),dpi=72),caption=caption)
    text,assets,_,_=normalize_docling_document(captioned)
    assert 'Actual caption' in text and assets[0]['alt']=='Actual caption'
